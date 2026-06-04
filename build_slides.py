#!/usr/bin/env python3
"""Build the PCD Feature Flags hackathon slide deck (pptx)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ----- Theme -----
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x60, 0x6E)
ACCENT = RGBColor(0x16, 0xA3, 0x4A)   # green
WARN = RGBColor(0xE6, 0x8A, 0x00)     # amber
DANGER = RGBColor(0xC0, 0x39, 0x2B)
BG = RGBColor(0xF7, 0xF8, 0xFA)
CODE_BG = RGBColor(0x1E, 0x29, 0x3B)
CODE_FG = RGBColor(0xE7, 0xEC, 0xF3)
SOFT = RGBColor(0xEA, 0xEE, 0xF3)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.shadow.inherit = False
    return slide


def add_text(slide, text, left, top, width, height,
             font_size=18, bold=False, color=INK,
             align=PP_ALIGN.LEFT, font_name="Helvetica Neue"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    if not isinstance(text, list):
        text = [text]
    for i, line in enumerate(text):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name
    return box


def add_bullets(slide, items, left, top, width, height,
                font_size=18, color=INK, bullet_color=NAVY, font_name="Helvetica Neue"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        # Bullet glyph
        b = p.add_run()
        b.text = "▸  "
        b.font.size = Pt(font_size)
        b.font.color.rgb = bullet_color
        b.font.bold = True
        b.font.name = font_name
        # Body
        r = p.add_run()
        r.text = item
        r.font.size = Pt(font_size)
        r.font.color.rgb = color
        r.font.name = font_name
    return box


def add_band(slide, top, height, color):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, SLIDE_W, height)
    band.line.fill.background()
    band.fill.solid()
    band.fill.fore_color.rgb = color
    band.shadow.inherit = False
    return band


def add_code(slide, code, left, top, width, height, font_size=14,
             fg=CODE_FG, bg=CODE_BG):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    rect.line.fill.background()
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.shadow.inherit = False
    rect.adjustments[0] = 0.04
    rect.text_frame.margin_left = Inches(0.2)
    rect.text_frame.margin_right = Inches(0.2)
    rect.text_frame.margin_top = Inches(0.15)
    rect.text_frame.margin_bottom = Inches(0.15)
    rect.text_frame.word_wrap = True
    lines = code.split("\n")
    for i, line in enumerate(lines):
        p = rect.text_frame.paragraphs[0] if i == 0 else rect.text_frame.add_paragraph()
        run = p.add_run()
        run.text = line if line else " "
        run.font.name = "Menlo"
        run.font.size = Pt(font_size)
        run.font.color.rgb = fg
    return rect


def add_pill(slide, text, left, top, width, height, color=ACCENT, fg=RGBColor(0xFF, 0xFF, 0xFF), font_size=16):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    rect.line.fill.background()
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.shadow.inherit = False
    rect.adjustments[0] = 0.5
    tf = rect.text_frame
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = fg
    run.font.name = "Helvetica Neue"
    return rect


def add_footer(slide, n, total):
    add_text(slide, f"PCD Feature Flags  ·  Hackathon 2026", Inches(0.5), Inches(7.1),
             Inches(6), Inches(0.3), font_size=10, color=MUTED)
    add_text(slide, f"{n} / {total}", Inches(12.3), Inches(7.1), Inches(0.8),
             Inches(0.3), font_size=10, color=MUTED, align=PP_ALIGN.RIGHT)


# ============== Slides ==============
TOTAL = 12

# 1 — Title
s = add_slide()
add_band(s, 0, Inches(7.5), NAVY)
add_text(s, "PCD Feature Flags", Inches(0.8), Inches(2.0),
         Inches(11.5), Inches(1.4), font_size=72, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))
add_text(s, "Flip a flag in 30 seconds.", Inches(0.8), Inches(3.6),
         Inches(11.5), Inches(0.7), font_size=36, bold=False,
         color=RGBColor(0xFF, 0xFF, 0xFF))
add_text(s, "No restart.  No redeploy.  No ticket.", Inches(0.8), Inches(4.3),
         Inches(11.5), Inches(0.7), font_size=28, color=ACCENT)
add_text(s, "mahij@platform9.com  ·  2026-06-04",
         Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.4),
         font_size=14, color=RGBColor(0xC0, 0xCB, 0xDC))

# 2 — The Problem
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "The Problem", Inches(0.5), Inches(0.15),
         Inches(12), Inches(0.5), font_size=28, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))

add_text(s, "PCD backend services have no runtime flag mechanism.",
         Inches(0.8), Inches(1.1), Inches(11.5), Inches(0.7),
         font_size=26, bold=True, color=INK)

add_bullets(s, [
    "The only flag system is /opt/pf9/www/public/clarity/features.json",
    "Ansible-driven · file on DU host · 12 years old · Clarity UI only",
    "Backend services (hamgr, nova, neutron, cinder, masakari, glance) can't gate behavior at runtime",
    "Every new code path ships either binary-on (scary) or behind static config (restart to toggle)",
], Inches(0.8), Inches(2.0), Inches(11.5), Inches(2.4), font_size=18)

add_text(s, "Meanwhile, DUs now run as k8s namespaces with helm-deployed pods.",
         Inches(0.8), Inches(4.7), Inches(11.5), Inches(0.6),
         font_size=20, color=MUTED)
add_text(s, "The 12-year-old features.json shape is obsolete.",
         Inches(0.8), Inches(5.3), Inches(11.5), Inches(0.6),
         font_size=22, bold=True, color=DANGER)
add_footer(s, 2, TOTAL)

# 3 — The Solution
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "The Solution", Inches(0.5), Inches(0.15),
         Inches(12), Inches(0.5), font_size=28, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))

add_text(s, "Adopt flagd — CNCF OpenFeature reference engine",
         Inches(0.8), Inches(1.1), Inches(11.5), Inches(0.7),
         font_size=26, bold=True, color=INK)

add_bullets(s, [
    "Stateless pod, one per DU namespace",
    "ConfigMap-backed · hot-reloads on kubectl edit",
    "Consumers use upstream OpenFeature SDK + flagd-provider directly",
    "No Platform9 wrapper.  No new pf9-flags repo.  ~20 LoC per service.",
], Inches(0.8), Inches(2.0), Inches(11.5), Inches(2.4), font_size=19)

add_pill(s, "Go consumer ready", Inches(0.8), Inches(5.0), Inches(2.8), Inches(0.5))
add_pill(s, "Python consumer ready", Inches(3.9), Inches(5.0), Inches(3.2), Inches(0.5))
add_pill(s, "Same flag schema for both", Inches(7.4), Inches(5.0), Inches(3.6), Inches(0.5),
         color=NAVY)

add_text(s, "Choose, integrate, ship — not invent.",
         Inches(0.8), Inches(6.0), Inches(11.5), Inches(0.6),
         font_size=20, bold=True, color=ACCENT)
add_footer(s, 3, TOTAL)

# 4 — Architecture
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "Architecture", Inches(0.5), Inches(0.15),
         Inches(12), Inches(0.5), font_size=28, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))

add_code(s, """\
+--------------- DU k8s namespace ---------------+
|                                                |
|   ConfigMap flagd-config                       |
|   { "flags": { "hamgr.fast_evac": {...} } }    |
|              |                                 |
|              | kubelet mounts as file (~30s)   |
|              v                                 |
|   flagd pod  --- gRPC :8013 push --->          |
|   (one per namespace)                          |
|              |              |             |    |
|              v              v             v    |
|         hamgr pod      nova-api      neutron   |
|        (OpenFeature   (OpenFeature   (future)  |
|         Go SDK)        Python SDK)             |
|                                                |
+------------------------------------------------+

    Operator:  kubectl edit cm/flagd-config -n <du-ns>
    Latency:   edit ─> service eval = ~30 seconds p99""",
         Inches(0.8), Inches(1.0), Inches(11.7), Inches(5.7), font_size=14)
add_footer(s, 4, TOTAL)

# 5 — What we shipped
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "What We Shipped — 7 PRs in 1 day, 3 repos",
         Inches(0.5), Inches(0.15), Inches(12.5), Inches(0.5),
         font_size=24, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

rows = [
    ("#677", "pf9-openstack-helm", "new flagd/ helm chart (foundation)"),
    ("#678", "pf9-openstack-helm", "kind-based e2e test"),
    ("#41",  "pf9-hamgr",          "Go consumer + gated hamgr.verbose_evac_logs"),
    ("#679", "pf9-openstack-helm", "hamgr chart env-var injection"),
    ("#410", "pf9-nova",           "Python consumer + gated nova.verbose_api_request_log"),
    ("#680", "pf9-openstack-helm", "nova chart env-var injection"),
    ("#681", "pf9-openstack-helm", "usage docs"),
]
row_top = Inches(1.1)
row_h = Inches(0.5)
add_text(s, "PR",   Inches(0.8), row_top, Inches(1.0), Inches(0.5),
         font_size=14, bold=True, color=MUTED)
add_text(s, "Repo", Inches(2.0), row_top, Inches(3.6), Inches(0.5),
         font_size=14, bold=True, color=MUTED)
add_text(s, "What it ships", Inches(5.8), row_top, Inches(7.0), Inches(0.5),
         font_size=14, bold=True, color=MUTED)
y = row_top + Inches(0.5)
for pr, repo, what in rows:
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y,
                              Inches(12.2), row_h)
    band.line.fill.background()
    band.fill.solid()
    band.fill.fore_color.rgb = SOFT
    band.shadow.inherit = False
    add_text(s, pr, Inches(0.8), y + Inches(0.1), Inches(1.0), Inches(0.4),
             font_size=16, bold=True, color=NAVY)
    add_text(s, repo, Inches(2.0), y + Inches(0.1), Inches(3.6), Inches(0.4),
             font_size=14, color=INK, font_name="Menlo")
    add_text(s, what, Inches(5.8), y + Inches(0.1), Inches(7.0), Inches(0.4),
             font_size=14, color=INK)
    y += row_h + Inches(0.08)

add_text(s, "Each PR has unit tests · helm-template snapshot tests · two-stage AI code review (spec compliance + code quality).",
         Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.6),
         font_size=14, color=MUTED)
add_footer(s, 5, TOTAL)

# 6 — Pilot consumers
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "Pilot Consumers (both languages, in parallel)",
         Inches(0.5), Inches(0.15), Inches(12.5), Inches(0.5),
         font_size=24, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

# Left column: Go / hamgr
add_text(s, "Go — pf9-hamgr", Inches(0.6), Inches(1.0),
         Inches(6), Inches(0.5), font_size=22, bold=True, color=NAVY)
add_code(s, """\
import "github.com/open-feature/go-sdk/openfeature"

func main() {
    initFlags()
    // ... existing startup ...
}

// At a call site:
client := openfeature.NewClient("hamgr")
on, _ := client.BooleanValue(ctx,
    "hamgr.fast_evac", false,
    openfeature.EvaluationContext{})
if on {
    fastEvacPath()
}""", Inches(0.5), Inches(1.5), Inches(6.2), Inches(5.0), font_size=13)

# Right column: Python / nova
add_text(s, "Python — pf9-nova", Inches(6.9), Inches(1.0),
         Inches(6), Inches(0.5), font_size=22, bold=True, color=NAVY)
add_code(s, """\
from nova.pf9 import flags
flags.init_flags()

# At a call site:
from openfeature import api
client = api.get_client('nova')
on = client.get_boolean_value(
    'nova.new_scheduler', False)
if on:
    new_scheduler_path()""",
         Inches(6.8), Inches(1.5), Inches(6.2), Inches(5.0), font_size=13)

add_text(s, "Same flag schema · same wire protocol · same SDK family.",
         Inches(0.6), Inches(6.7), Inches(12), Inches(0.5),
         font_size=16, color=MUTED, align=PP_ALIGN.CENTER)
add_footer(s, 6, TOTAL)

# 7 — DEMO TIME
s = add_slide()
add_band(s, 0, Inches(7.5), NAVY)
add_text(s, "DEMO", Inches(0), Inches(2.3), SLIDE_W, Inches(1.5),
         font_size=120, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
         align=PP_ALIGN.CENTER)
add_text(s, "3 terminals · 3 commands · 1 minute", Inches(0), Inches(4.2),
         SLIDE_W, Inches(0.8), font_size=28, color=ACCENT,
         align=PP_ALIGN.CENTER)
add_code(s, """\
T1 $ ./start-flagd.sh        # flagd boots, watches flags.json
T2 $ ./start-service.sh      # fake hamgr starts polling
T3 $ ./flip.sh on            # service flips to green FAST path in 1s
T3 $ ./flip.sh off           # service rolls back in 1s""",
         Inches(1.5), Inches(5.2), Inches(10.3), Inches(1.7), font_size=14,
         fg=RGBColor(0xE7, 0xEC, 0xF3), bg=RGBColor(0x00, 0x10, 0x25))
add_footer(s, 7, TOTAL)

# 8 — What the audience saw (transcript)
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "What the audience just saw", Inches(0.5), Inches(0.15),
         Inches(12), Inches(0.5), font_size=24, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))

add_code(s, """\
[19:53:08] cycle=1   legacy evac path: serial-migrating 8 VMs in 96s
[19:53:09] cycle=2   legacy evac path: serial-migrating 8 VMs in 96s
                                                                       <-- ./flip.sh on
[19:53:10] cycle=3   FAST evac path: parallel-migrating 8 VMs in 12s
            ↑ flag flipped ON — switched to fast path on cycle 3
[19:53:11] cycle=4   FAST evac path: parallel-migrating 8 VMs in 12s
[19:53:12] cycle=5   FAST evac path: parallel-migrating 8 VMs in 12s
                                                                       <-- ./flip.sh off
[19:53:15] cycle=8   legacy evac path: serial-migrating 8 VMs in 96s
            ↓ flag flipped OFF — rolled back to legacy path on cycle 8
[19:53:16] cycle=9   legacy evac path: serial-migrating 8 VMs in 96s""",
         Inches(0.5), Inches(1.0), Inches(12.3), Inches(5.6), font_size=14)
add_text(s, "Effect propagation: ~1 second.  No restart.  No redeploy.",
         Inches(0.5), Inches(6.7), Inches(12), Inches(0.5),
         font_size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
add_footer(s, 8, TOTAL)

# 9 — Production parity
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "Production Parity — demo IS the prod path",
         Inches(0.5), Inches(0.15), Inches(12.5), Inches(0.5),
         font_size=24, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

rows = [
    ("Demo (laptop)", "Production (DU)"),
    ("flags.json on local disk", "ConfigMap flagd-config mounted at /etc/flagd/"),
    ("flagd as a host binary", "flagd as a k8s Deployment, one per DU namespace"),
    ("Service runs against localhost:8013", "Pod injects FLAGD_HOST / FLAGD_PORT via helm"),
    ("./flip.sh on", "kubectl edit cm/flagd-config -n <du-ns>"),
    ("Same OpenFeature SDK", "Same OpenFeature SDK"),
    ("Same flagd-provider", "Same flagd-provider"),
    ("Same JSON flag schema", "Same JSON flag schema"),
]
row_top = Inches(1.1)
row_h = Inches(0.55)
for i, (left, right) in enumerate(rows):
    bg_color = NAVY if i == 0 else (SOFT if i % 2 else BG)
    fg_color = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
    bold = i == 0
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5),
                              row_top + i * (row_h + Inches(0.02)),
                              Inches(12.3), row_h)
    rect.line.fill.background()
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg_color
    rect.shadow.inherit = False
    add_text(s, left,  Inches(0.7),
             row_top + i * (row_h + Inches(0.02)) + Inches(0.1),
             Inches(5.5), Inches(0.4),
             font_size=15, bold=bold, color=fg_color)
    add_text(s, right, Inches(6.5),
             row_top + i * (row_h + Inches(0.02)) + Inches(0.1),
             Inches(6.5), Inches(0.4),
             font_size=15, bold=bold, color=fg_color)
add_footer(s, 9, TOTAL)

# 10 — Why this wins
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "Why this wins", Inches(0.5), Inches(0.15),
         Inches(12), Inches(0.5), font_size=28, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))

add_bullets(s, [
    "Zero Platform9 reinvention — flagd is CNCF, OpenFeature is the standard",
    "Already production-ready — unit tests + snapshot tests + e2e + AI code review on every PR",
    "Working code, not slides — the demo runs on this laptop in 60 seconds",
    "Extensible — same pattern adds neutron / cinder / glance in <1 day each",
    "Operator UX — kubectl edit cm — no new tool to learn, K8s RBAC already gates who can flip",
    "Fail-open — every consumer degrades to caller default on flagd outage; service never crashes",
], Inches(0.8), Inches(1.2), Inches(12), Inches(5.5), font_size=19)
add_footer(s, 10, TOTAL)

# 11 — What's next (v2)
s = add_slide()
add_band(s, 0, Inches(0.7), NAVY)
add_text(s, "What's next  (v2 candidates)",
         Inches(0.5), Inches(0.15), Inches(12), Inches(0.5),
         font_size=28, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

add_text(s, "Each line is a candidate for a follow-up spec when concrete demand arises:",
         Inches(0.8), Inches(1.1), Inches(11.5), Inches(0.5),
         font_size=16, color=MUTED)

add_bullets(s, [
    "Per-tenant / per-project / per-user targeting (already in flagd's flag schema)",
    "Percent rollouts and experimentation",
    "Hypervisor-side services (nova-compute, pf9-ha-slave, vmha-agent on bare metal)",
    "String / int / JSON flag types beyond bool",
    "PCD CLI wrapper around kubectl edit (if operator UX feedback warrants)",
    "Audit dashboard beyond K8s native audit log",
], Inches(0.8), Inches(2.0), Inches(12), Inches(4.5), font_size=19)

add_pill(s, "All deferred features are flagd-native — zero consumer-side rewrite needed",
         Inches(0.8), Inches(6.4), Inches(11.5), Inches(0.5),
         color=ACCENT, font_size=14)
add_footer(s, 11, TOTAL)

# 12 — Thanks / closing
s = add_slide()
add_band(s, 0, Inches(7.5), NAVY)
add_text(s, "Thanks.", Inches(0.8), Inches(2.0),
         Inches(11.5), Inches(1.6), font_size=96, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))
add_text(s, "No restart.  No redeploy.  No incident.",
         Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.8),
         font_size=34, color=ACCENT)
add_text(s, "Code:  /tmp/pcd-flags-demo  ·  spec, plan, 7 PRs all linked in README.md",
         Inches(0.8), Inches(5.6), Inches(11.5), Inches(0.5),
         font_size=14, color=RGBColor(0xC0, 0xCB, 0xDC))
add_text(s, "mahij@platform9.com",
         Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.4),
         font_size=14, color=RGBColor(0xC0, 0xCB, 0xDC))

out = "/tmp/pcd-flags-demo/pcd-feature-flags.pptx"
prs.save(out)
print(f"saved: {out}")
